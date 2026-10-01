"""
Contrôle + export de vérification.
    python verifier.py            -> affiche les anomalies et crée export_verification.xlsx
Code retour 1 s'il reste une anomalie BLOQUANTE (utile avant chaque publication).
"""
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

import data
import controles

SORTIE = "export_verification.xlsx"


def main() -> int:
    base = data.charger()
    bl = [a for a in base.anomalies if a.niveau == controles.BLOQUANT]
    al = [a for a in base.anomalies if a.niveau == controles.ALERTE]
    print(f"{len(base.buts)} buts · {len(base.matchs)} matchs · journées {base.journees}")
    print(f"{len(bl)} bloquant(s) · {len(al)} alerte(s)")
    for a in bl + al:
        print(" ", a)

    s = base.stats_joueurs()
    lig = base.stats_ligue()
    clubs = []
    for c in base.clubs.club_court:
        sc = base.stats_club(c)
        row = {"club": c, **{k: sc["classement"][k] for k in ["Rg", "Pts", "J", "G", "N", "P", "BP", "BC", "Diff"]}}
        dom = base.classement("dom").set_index("club").loc[c]; ext = base.classement("ext").set_index("club").loc[c]
        row.update({"Pts_dom": dom.Pts, "Pts_ext": ext.Pts, "BP_1MT": int(((base.buts.club_marque == c) & (base.buts.periode == 1)).sum()),
                    "BP_2MT": int(((base.buts.club_marque == c) & (base.buts.periode == 2)).sum())})
        row.update({f"orig_{k}": v for k, v in sc["origines"].items()})
        clubs.append(row)
    b = base.buts[["journee", "club_dom", "club_ext", "periode", "minute", "club_marque", "buteur", "passeur",
                   "origine_but", "score_dom_apres", "score_ext_apres"]]
    m = base.matchs.assign(score=lambda d: d.bd.astype(str) + "-" + d.be.astype(str),
                           mi_temps=lambda d: d.mt_d.astype(str) + "-" + d.mt_e.astype(str))[["journee", "dom", "score", "ext", "mi_temps", "nb_buts"]]
    ctrl = pd.DataFrame([{"niveau": a.niveau, "source": a.source, "message": a.message} for a in bl + al]) \
        if (bl or al) else pd.DataFrame([{"niveau": "OK", "source": "", "message": "Aucune anomalie."}])
    ligue = pd.DataFrame([{"indicateur": k, "valeur": v} for k, v in lig.items() if not isinstance(v, dict)]
                         + [{"indicateur": f"origine · {k}", "valeur": v} for k, v in lig["origines"].items()]
                         + [{"indicateur": f"tranche {k}'", "valeur": v} for k, v in lig["tranches"].items()])

    feuilles = {"Contrôles": ctrl, "Classement": base.classement(), "Buteurs": base.buteurs(),
                "Passeurs": base.passeurs(), "Buts+Passes": base.buts_plus_passes(),
                "Joueurs (complet)": s.drop(columns=["id_joueur"]), "Clubs": pd.DataFrame(clubs).fillna(0),
                "Matchs": m, "Buts": b, "Ligue": ligue}
    with pd.ExcelWriter(SORTIE, engine="openpyxl") as w:
        for nom, df in feuilles.items():
            df.to_excel(w, sheet_name=nom, index=False)

    # mise en forme lisible
    wb = load_workbook(SORTIE)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="111111")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial")
        for col in ws.columns:
            larg = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(8, larg + 2), 60 if ws.title != "Contrôles" else 120)
        ws.freeze_panes = "A2"
    ws = wb["Contrôles"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "BLOQUANT":
            row[0].font = Font(name="Arial", bold=True, color="C8102E")
    wb.save(SORTIE)
    print(f"Export : {SORTIE}")
    return 1 if bl else 0


if __name__ == "__main__":
    sys.exit(main())
