"""
Générateur de textes éditoriaux (La Une, l'essentiel de chaque match).

Chaque phrase est construite à partir de règles factuelles vérifiables dans les données
(score, renversement, doublé, but tardif, phase de jeu dominante, record...).
Rien n'est publié sans validation : le site n'affiche que les lignes au statut OK.

Usage :
    python generer_textes.py          -> propose les textes de la dernière journée
    python generer_textes.py 4        -> journée 4
Le fichier textes_journee.xlsx est mis à jour sans écraser ce qui est déjà validé.
"""
from __future__ import annotations
import re
import sys
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

import config as C
import data

COLONNES = ["cible", "journee", "type", "match", "texte_propose", "texte_valide", "statut", "faits"]
PHASES = {"attaque placée": "en attaque placée", "attaque rapide": "en attaque rapide", "transition off": "en transition",
          "corner": "sur corner", "touche off": "sur touche", "touche def": "sur touche défensive",
          "coup franc": "sur coup franc", "jet franc": "sur jet franc", "power play": "en power play", "penalty": "sur penalty"}
NOMBRES = {2: "deux", 3: "trois", 4: "quatre", 5: "cinq", 6: "six", 7: "sept", 8: "huit"}


def _nom(base, club):
    r = base.clubs.set_index("club_court")
    v = r.nom_affiche.get(club) if "nom_affiche" in r else None
    return v if isinstance(v, str) and v else C.NOMS_CLUBS.get(club, club)


def _court(nom):
    """'Francisco Martinez Villalba' -> 'Martinez Villalba' (on retire le prénom)."""
    p = nom.split()
    return " ".join(p[1:]) if len(p) > 1 else nom


def _n(x):
    return NOMBRES.get(x, str(x))


# ------------------------------------------------------------------ faits d'un match
def faits_match(base, m) -> dict:
    """Tout ce qu'on peut affirmer sur un match, calculé depuis les buts."""
    b = base.match(m.journee, m.dom)
    f = {"dom": m.dom, "ext": m.ext, "bd": m.bd, "be": m.be, "nb": len(b)}
    f["vainqueur"] = m.dom if m.bd > m.be else (m.ext if m.be > m.bd else None)
    f["perdant"] = m.ext if f["vainqueur"] == m.dom else (m.dom if f["vainqueur"] == m.ext else None)
    f["ecart"] = abs(m.bd - m.be)
    f["mi_temps"] = (m.mt_d, m.mt_e)
    # renversement : le vainqueur a été mené à un moment
    if f["vainqueur"]:
        v_dom = f["vainqueur"] == m.dom
        retard = [(r.score_ext_apres - r.score_dom_apres) if v_dom else (r.score_dom_apres - r.score_ext_apres) for r in b.itertuples()]
        f["mene_max"] = max([0] + retard)
        f["mene_mt"] = (m.mt_e > m.mt_d) if v_dom else (m.mt_d > m.mt_e)
        # avance maximale du vainqueur et retour du perdant
        avance = [-x for x in retard]
        f["avance_max"] = max([0] + avance)
        # but décisif : celui qui donne l'avance définitive
        dec = None
        for r in b.itertuples():
            d = (r.score_dom_apres - r.score_ext_apres) if v_dom else (r.score_ext_apres - r.score_dom_apres)
            if d >= 1 and r.club_marque == f["vainqueur"] and (dec is None or prev_d <= 0):
                dec = r
            prev_d = d
        f["decisif"] = dec
    # premier but
    f["premier"] = b.iloc[0] if len(b) else None
    # buteurs multiples
    bt = b[b.id_buteur.notna()].groupby(["buteur", "club_marque"]).size().sort_values(ascending=False)
    f["multi"] = [(n, c, k) for (n, c), k in bt.items() if k >= 2]
    # phase dominante par équipe (au moins 3 buts et la moitié)
    f["phase"] = {}
    for club in (m.dom, m.ext):
        o = b[b.club_marque == club].origine_but.value_counts()
        if len(o) and o.iloc[0] >= 3 and o.iloc[0] * 2 >= o.sum():
            f["phase"][club] = (o.index[0], int(o.iloc[0]), int(o.sum()))
    # record de buts sur la saison jusqu'à cette journée
    avant = base.matchs[base.matchs.journee <= m.journee]
    f["record"] = len(b) == avant.nb_buts.max() and (avant.nb_buts == len(b)).sum() == 1 and m.journee > 1
    f["blanchi"] = f["perdant"] if f["vainqueur"] and min(m.bd, m.be) == 0 else None
    f["passeur_top"] = None
    pt = b[b.id_passeur.notna()].groupby(["passeur", "club_marque"]).size().sort_values(ascending=False)
    if len(pt) and pt.iloc[0] >= 2:
        f["passeur_top"] = (pt.index[0][0], pt.index[0][1], int(pt.iloc[0]))
    return f


# ------------------------------------------------------------------ grammaire
VOYELLES = "AEIOUYÉÈÊÂÎÔH"


def _art(base, c):
    """('le Sporting', False) / ('Laval', False) / ('les Artistes', True)."""
    if c in C.ARTICLES_CLUBS:
        return C.ARTICLES_CLUBS[c]
    return _nom(base, c), False


def sujet(base, c):
    t = _art(base, c)[0]
    return t[0].upper() + t[1:]


def objet(base, c):
    return _art(base, c)[0]


def de(texte):
    """de + nom : du Sporting, des Artistes, d'Avion, de Laval."""
    if texte.startswith("le "):
        return "du " + texte[3:]
    if texte.startswith("les "):
        return "des " + texte[4:]
    return ("d’" if texte[:1].upper() in VOYELLES else "de ") + texte


def pl(base, c):
    return _art(base, c)[1]


def verbe(base, c, sing, plur):
    return plur if pl(base, c) else sing


# ------------------------------------------------------------------ phrases
def essentiel(base, f) -> tuple[str, list[str]]:
    """Une phrase (deux au plus) qui résume le match. Renvoie (texte, faits utilisés)."""
    S = lambda c: sujet(base, c)
    O = lambda c: objet(base, c)
    V = lambda c, a, b: verbe(base, c, a, b)
    faits = []
    v, p = f["vainqueur"], f["perdant"]
    if v is None:
        phrase = f"{S(f['dom'])} et {O(f['ext'])} se quittent sur un nul {f['bd']}-{f['be']}"
        faits.append("nul")
    else:
        sv, sp = max(f["bd"], f["be"]), min(f["bd"], f["be"])
        mene = V(v, "Mené", "Menés")
        if f["mene_max"] >= 2:
            phrase = f"{mene} de {_n(f['mene_max'])} buts, {O(v)} {V(v, 'renverse', 'renversent')} {O(p)} ({sv}-{sp})"
            faits.append(f"renversement après {f['mene_max']} buts de retard")
        elif f["mene_mt"]:
            phrase = f"{mene} à la pause, {O(v)} {V(v, 'renverse', 'renversent')} {O(p)} ({sv}-{sp})"
            faits.append("mené à la mi-temps")
        elif f["avance_max"] >= 4 and f["ecart"] <= 2:
            phrase = (f"{S(v)} {V(v, 's’impose', 's’imposent')} {sv}-{sp} contre {O(p)} mais {V(v, 'a', 'ont')} tremblé "
                      f"après avoir mené de {_n(f['avance_max'])} buts")
            faits.append(f"avance de {f['avance_max']} réduite à {f['ecart']}")
        elif f["blanchi"]:
            phrase = (f"{S(v)} {V(v, 'écarte', 'écartent')} {O(p)} {sv}-0 sans encaisser" if f["ecart"] >= 4
                      else f"{S(v)} {V(v, 'bat', 'battent')} {O(p)} {sv}-0 sans encaisser")
            faits.append("victoire sans but encaissé")
        elif f["ecart"] >= 4:
            phrase = f"Large succès {de(O(v))} contre {O(p)} ({sv}-{sp})"
            faits.append(f"écart de {f['ecart']}")
        elif f["ecart"] == 1 and f["decisif"] is not None and f["decisif"].minute >= 35:
            d = f["decisif"]
            phrase = (f"{S(v)} {V(v, 'arrache', 'arrachent')} la victoire face à {O(p)} ({sv}-{sp}) "
                      f"sur un but {de(_court(d.buteur))} à la {d.minute}e")
            faits.append(f"but décisif à la {d.minute}e")
        else:
            phrase = f"{S(v)} {V(v, 'bat', 'battent')} {O(p)} {sv}-{sp}"
            faits.append("victoire")
    # un seul détail : doublé/triplé, sinon phase dominante, sinon passeur
    if f["multi"]:
        n, c, k = f["multi"][0]
        mot = {2: "doublé", 3: "triplé", 4: "quadruplé"}.get(k, f"total de {k} buts")
        phrase += f", avec un {mot} {de(_court(n))}"
        faits.append(f"{n} {k} buts")
    elif v and v in f["phase"]:
        ph, k, tot = f["phase"][v]
        phrase += f", avec {_n(k)} de ses {_n(tot)} buts {PHASES.get(ph, ph)}"
        faits.append(f"{k}/{tot} buts {ph}")
    elif f["passeur_top"]:
        n, c, k = f["passeur_top"]
        phrase += f", avec {_n(k)} passes décisives {de(_court(n))}"
        faits.append(f"{n} {k} passes")
    texte = phrase + "."
    if f["record"]:
        texte += f" Avec {f['nb']} buts, c’est le match le plus prolifique de la saison."
        faits.append(f"record {f['nb']} buts")
    return texte, faits


def titre_une(base, f) -> tuple[str, list[str]]:
    """Titre court pour La Une (match de la journée)."""
    S = lambda c: sujet(base, c)
    O = lambda c: objet(base, c)
    V = lambda c, a, b: verbe(base, c, a, b)
    v, p = f["vainqueur"], f["perdant"]
    if f["record"]:
        return f"{f['nb']} buts entre {O(f['dom'])} et {O(f['ext'])}, record de la saison", [f"record {f['nb']} buts"]
    if v and f["mene_max"] >= 2:
        return f"{S(v)} {V(v, 'renverse', 'renversent')} {O(p)}", ["renversement"]
    if f["multi"] and f["multi"][0][2] >= 3:
        n, c, k = f["multi"][0]
        return f"{_court(n)} signe un {'triplé' if k == 3 else 'quadruplé' if k == 4 else f'total de {k} buts'}", [f"{n} {k} buts"]
    if v and f["phase"].get(v) and f["phase"][v][1] >= 4:
        ph, k, tot = f["phase"][v]
        return f"{S(v)} {V(v, 'fait', 'font')} la différence {PHASES.get(ph, ph)}", [f"{k}/{tot} {ph}"]
    if f["nb"] >= 9:
        return f"{f['nb']} buts entre {O(f['dom'])} et {O(f['ext'])}", [f"{f['nb']} buts"]
    if v and f["ecart"] >= 4:
        return f"{S(v)} {V(v, 'déroule', 'déroulent')} face à {O(p)}", [f"écart {f['ecart']}"]
    if v:
        return f"{S(v)} {V(v, 'prend', 'prennent')} le dessus sur {O(p)}", ["victoire"]
    return f"{S(f['dom'])} et {O(f['ext'])} dos à dos", ["nul"]


# ------------------------------------------------------------------ fichier
def match_une(base, j):
    """Même choix que l'appli : le match le plus riche en buts de la journée."""
    m = base.matchs[base.matchs.journee == j]
    return m.sort_values(["nb_buts", "dom"], ascending=[False, True]).iloc[0]


def propositions(base, j) -> pd.DataFrame:
    rows = []
    top = match_une(base, j)
    for m in base.matchs[base.matchs.journee == j].itertuples():
        f = faits_match(base, m)
        etiquette = f"{_nom(base, m.dom)} {m.bd}-{m.be} {_nom(base, m.ext)}"
        if m.dom == top.dom:
            t, fa = titre_une(base, f)
            rows.append(dict(cible=f"une:J{j}", journee=j, type="Titre de La Une", match=etiquette, texte_propose=t, faits=" ; ".join(fa)))
        t, fa = essentiel(base, f)
        rows.append(dict(cible=f"match:J{j}:{m.dom}", journee=j, type="Essentiel du match", match=etiquette, texte_propose=t, faits=" ; ".join(fa)))
    df = pd.DataFrame(rows)
    df["texte_propose"] = df.texte_propose.map(_contracter)
    return df


def _contracter(t: str) -> str:
    """« face à le Sporting » -> « face au Sporting », « de les Artistes » -> « des Artistes »."""
    for a, b in (("à le ", "au "), ("à les ", "aux "), ("de le ", "du "), ("de les ", "des ")):
        t = re.sub(rf"\b{a}", b, t)
    return t


def fusionner(existant: pd.DataFrame | None, nouveau: pd.DataFrame) -> pd.DataFrame:
    """Garde tout ce qui est déjà validé ; met à jour les propositions non validées."""
    if existant is None or existant.empty:
        out = nouveau.copy()
        out["texte_valide"], out["statut"] = "", ""
        return out[COLONNES]
    ex = existant.copy()
    for c in COLONNES:
        if c not in ex:
            ex[c] = ""
    ex = ex.set_index("cible")
    for r in nouveau.itertuples():
        if r.cible in ex.index:
            if str(ex.at[r.cible, "statut"]).strip().upper() != "OK":
                for c in ["journee", "type", "match", "texte_propose", "faits"]:
                    ex.at[r.cible, c] = getattr(r, c)
        else:
            ex.loc[r.cible] = {**{c: getattr(r, c) for c in ["journee", "type", "match", "texte_propose", "faits"]}, "texte_valide": "", "statut": ""}
    out = ex.reset_index()[COLONNES]
    return out.sort_values(["journee", "type", "cible"], ascending=[False, False, True]).reset_index(drop=True)


def lire_existant(path=C.FICHIER_TEXTES) -> pd.DataFrame | None:
    return pd.read_excel(path).fillna("") if path.exists() else None


def ecrire(df: pd.DataFrame, cible) -> None:
    """Écrit le fichier mis en forme (chemin ou BytesIO)."""
    df.to_excel(cible, index=False, sheet_name="Textes")
    if isinstance(cible, BytesIO):
        cible.seek(0)
    wb = load_workbook(cible)
    ws = wb["Textes"]
    larg = {"A": 20, "B": 8, "C": 18, "D": 30, "E": 70, "F": 70, "G": 10, "H": 40}
    for col, w in larg.items():
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="111111")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", color="888888" if c.column_letter in "ABCH" else "111111")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[5].fill = PatternFill("solid", fgColor="FBF8F2")
    dv = DataValidation(type="list", formula1='"OK,A REVOIR"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{max(ws.max_row, 2) + 50}")
    ws.freeze_panes = "E2"
    if isinstance(cible, BytesIO):
        cible.seek(0); cible.truncate()
    wb.save(cible)


def fichier_pour_telechargement(base, j) -> bytes:
    buf = BytesIO()
    ecrire(fusionner(lire_existant(), propositions(base, j)), buf)
    return buf.getvalue()


if __name__ == "__main__":
    base = data.charger()
    if base.bloquant:
        sys.exit("Contrôle BLOQUANT en cours : corriger les données avant de générer les textes (python verifier.py).")
    j = int(sys.argv[1]) if len(sys.argv) > 1 else max(base.journees)
    df = fusionner(lire_existant(), propositions(base, j))
    ecrire(df, C.FICHIER_TEXTES)
    print(f"{C.FICHIER_TEXTES.name} : {len(df[df.journee == j])} textes proposés pour la J{j}. "
          "Relire, corriger dans texte_valide si besoin, mettre OK dans statut.")
